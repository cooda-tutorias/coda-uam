from datetime import date
from io import BytesIO
from types import SimpleNamespace

import docx
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.test import SimpleTestCase

from .services.cartas_asignacion import cargar_plantilla, generar_carta, validar_plantilla, parrafos
from .services.oficios import normalizar_numero_oficio


class CartasAsignacionTests(SimpleTestCase):
    def setUp(self):
        self.tutor = SimpleNamespace(matricula='99', first_name='Ana', last_name='Pérez', second_last_name='López', sexo='F',
                                     get_coordinacion_display=lambda: 'Ingeniería en Computación')
        self.alumno = SimpleNamespace(matricula='123', first_name='Luis', last_name='Ruiz', second_last_name='',
                                      trimestre_ingreso='26-O', get_carrera_display=lambda: 'Ingeniería en Computación')

    def plantilla(self, tipo='alumno'):
        documento = docx.Document()
        p = documento.add_paragraph('Oficio No. ')
        p.add_run('{no_').bold = True
        p.add_run('oficio}')
        documento.add_paragraph('{fecha}')
        if tipo == 'alumno':
            documento.add_paragraph('{nombre_alumno} de matrícula {matricula}')
            documento.add_paragraph('{asignado} {tutor} {academico} {el} {dr} {nombre_tutor} {licenciatura}')
        else:
            documento.add_paragraph('{nombre_mayus_tutor} {licenciatura}')
            tabla = documento.add_table(rows=2, cols=5)
            tabla.rows[1].cells[0].text = 'ALUMNO ANTERIOR'
        return documento

    def bytes(self, documento):
        salida = BytesIO()
        documento.save(salida)
        return salida.getvalue()

    def texto(self, contenido):
        return '\n'.join(p.text for p in parrafos(docx.Document(BytesIO(contenido))))

    def test_formato_compartido_y_anio_emision(self):
        self.assertEqual(normalizar_numero_oficio('003', date(2025, 1, 1)), 'DCNI_CODDAA_3_2025')

    def test_carta_individual_reemplaza_matricula_y_concordancia(self):
        contenido = generar_carta(self.bytes(self.plantilla()), self.tutor, [self.alumno], 63, date(2026, 9, 15), 'alumno')
        texto = self.texto(contenido)
        self.assertIn('DCNI_CODDAA_63_2026', texto)
        self.assertIn('de matrícula 123', texto)
        self.assertIn('asignada tutora académica la Dra. Ana Pérez López', texto)
        self.assertNotIn('{', texto)

    def test_carta_tutor_solo_contiene_alumnos_proporcionados(self):
        contenido = generar_carta(self.bytes(self.plantilla('tutor')), self.tutor, [self.alumno], 63, date(2026, 9, 15), 'tutor')
        documento = docx.Document(BytesIO(contenido))
        self.assertEqual(len(documento.tables[0].rows), 2)
        self.assertEqual(documento.tables[0].rows[1].cells[1].text, '123')
        self.assertNotIn('ALUMNO ANTERIOR', self.texto(contenido))

    def test_marcadores_faltantes_y_obsoletos(self):
        documento = self.plantilla()
        documento.paragraphs[2].text = '{nombre_alumno} {anio}'
        with self.assertRaises(ValidationError) as error:
            validar_plantilla(documento, 'alumno')
        self.assertIn('{matricula}', str(error.exception))
        self.assertIn('{anio}', str(error.exception))

    def test_tabla_incompatible(self):
        documento = self.plantilla()
        documento.add_table(rows=1, cols=1)
        with self.assertRaisesMessage(ValidationError, 'no debe contener tablas'):
            validar_plantilla(documento, 'alumno')

    def test_prefijo_duplicado(self):
        documento = self.plantilla()
        documento.paragraphs[0].text = 'DCNI_CODDAA_{no_oficio}'
        with self.assertRaisesMessage(ValidationError, 'Elimina el prefijo'):
            validar_plantilla(documento, 'alumno')

    def test_archivo_invalido(self):
        with self.assertRaisesMessage(ValidationError, '.docx válido'):
            cargar_plantilla(ContentFile(b'no es un word', name='archivo.docx'), 'alumno')




    def test_ejemplos_adaptados(self):
        from pathlib import Path
        directorio = Path(__file__).resolve().parent.parent / 'plantillas_ejemplo'
        for tipo in ('alumno', 'tutor'):
            contenido = (directorio / f'carta_asignacion_{tipo}_plantilla.docx').read_bytes()
            resultado = generar_carta(contenido, self.tutor, [self.alumno], 63, date(2026, 9, 15), tipo)
            self.assertNotIn('{', self.texto(resultado))

    def test_reporte_tutorias_conserva_formato_de_oficio(self):
        from .services.docx_reportes import _replace_reporte_placeholders
        documento = docx.Document()
        documento.add_paragraph('Oficio No. {no_oficio}')
        _replace_reporte_placeholders(documento, self.tutor,
                                     normalizar_numero_oficio(63, date(2026, 9, 15)), '2026-09-15T12:00')
        self.assertEqual(documento.paragraphs[0].text, 'Oficio No. DCNI_CODDAA_63_2026')



"""Validación y generación de cartas de asignación, sin respuestas HTTP."""
from copy import deepcopy
from io import BytesIO
import re
from zipfile import BadZipFile

import docx
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from django.core.exceptions import ValidationError

from .docx_reportes import paragraph_replace_text
from .oficios import normalizar_numero_oficio

OBLIGATORIOS = {
    'reporte': {'no_oficio', 'fecha', 'nombre_mayus_tutor', 'nombre_tutor'},
    'alumno': {'no_oficio', 'fecha', 'nombre_alumno', 'matricula', 'nombre_tutor'},
    'tutor': {'no_oficio', 'fecha', 'nombre_mayus_tutor', 'licenciatura'},
}
OPCIONALES = {
    'reporte': {'estimado'},
    'alumno': {'licenciatura', 'tutor', 'academico', 'dr', 'el', 'la', 'él', 'asignado', 'reasignado', 'a', 'o', 'e'},
    'tutor': {'nombre_tutor', 'estimado'},
}


def parrafos(documento):
    """Incluye tablas, encabezados y pies, también con tablas anidadas."""
    def recorrer(contenedor):
        yield from contenedor.paragraphs
        for tabla in contenedor.tables:
            for fila in tabla.rows:
                for celda in fila.cells:
                    yield from recorrer(celda)
    yield from recorrer(documento)
    for section in documento.sections:
        for parte in (section.header, section.first_page_header, section.even_page_header,
                      section.footer, section.first_page_footer, section.even_page_footer):
            yield from recorrer(parte)


def cargar_plantilla(archivo, tipo):
    try:
        archivo.open('rb')
        try:
            contenido = archivo.read()
        finally:
            archivo.close()
        documento = docx.Document(BytesIO(contenido))
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, XMLSyntaxError) as error:
        raise ValidationError('No se pudo abrir la plantilla. Selecciona un archivo Word .docx válido y disponible.') from error
    validar_plantilla(documento, tipo)
    return contenido


def validar_plantilla(documento, tipo):
    textos = [p.text for p in parrafos(documento)]
    marcadores = set(re.findall(r'\{([^{}]+)\}', '\n'.join(textos)))
    errores = []
    faltantes = OBLIGATORIOS[tipo] - marcadores
    desconocidos = marcadores - OBLIGATORIOS[tipo] - OPCIONALES[tipo]
    if faltantes:
        errores.append('Faltan marcadores obligatorios: ' + ', '.join('{' + m + '}' for m in sorted(faltantes)) + '.')
    if desconocidos:
        errores.append('Marcadores no admitidos: ' + ', '.join('{' + m + '}' for m in sorted(desconocidos)) + '. Usa «Oficio No. {no_oficio}»; el oficio ya incluye el año.')
    if any(re.search(r'DCNI\s*_\s*CODDAA\s*_\s*\{no_oficio\}', texto) for texto in textos):
        errores.append('Elimina el prefijo DCNI_CODDAA_ de la plantilla: usa «Oficio No. {no_oficio}».')
    if tipo == 'alumno' and documento.tables:
        errores.append('La carta para el alumno no debe contener tablas. Selecciona la plantilla dirigida al alumno.')
    if tipo == 'reporte' and not documento.tables:
        errores.append('El reporte de tutorías necesita una tabla para insertar las tutorías.')
    if tipo == 'tutor':
        if not documento.tables or len(documento.tables[0].rows) < 1 or any(len(f.cells) != 5 for f in documento.tables[0].rows):
            errores.append('La carta para el tutor necesita una tabla de cinco columnas: trimestre, matrícula, primer apellido, segundo apellido y nombres.')
    if errores:
        raise ValidationError(errores)


def generar_carta(plantilla, tutor, alumnos, oficio, fecha, tipo, licenciatura=None):
    documento = docx.Document(BytesIO(plantilla))
    validar_plantilla(documento, tipo)
    nombre = ' '.join(filter(None, [tutor.first_name, tutor.last_name, tutor.second_last_name]))
    femenino = tutor.sexo == 'F'
    tratamiento = 'Dra.' if femenino else 'Dr.'
    valores = {
        'no_oficio': normalizar_numero_oficio(oficio, fecha), 'fecha': fecha.isoformat(),
        'nombre_tutor': nombre, 'nombre_mayus_tutor': f'{tratamiento} {nombre}'.upper(),
        'estimado': 'Estimada' if femenino else 'Estimado',
        'licenciatura': licenciatura or tutor.get_coordinacion_display(),
        'tutor': 'tutora' if femenino else 'tutor', 'academico': 'académica' if femenino else 'académico',
        'dr': tratamiento, 'el': 'la' if femenino else 'el', 'la': 'la' if femenino else 'el',
        'él': 'ella' if femenino else 'él', 'asignado': 'asignada' if femenino else 'asignado',
        'reasignado': 'reasignada' if femenino else 'reasignado',
        'a': 'a' if femenino else 'o', 'o': 'a' if femenino else 'o', 'e': 'a' if femenino else 'e',
    }
    if tipo == 'alumno':
        alumno = alumnos[0]
        valores.update(nombre_alumno=' '.join(filter(None, [alumno.first_name, alumno.last_name, alumno.second_last_name])).upper(),
                       matricula=str(alumno.matricula), licenciatura=alumno.get_carrera_display().upper())
    else:
        valores['nombre_tutor'] = f'{tratamiento} {nombre}'
        tabla = documento.tables[0]
        modelo = deepcopy(tabla.rows[1]._tr) if len(tabla.rows) > 1 else None
        for fila in list(tabla.rows)[1:]:
            tabla._tbl.remove(fila._tr)
        for alumno in alumnos:
            if modelo is not None:
                tabla._tbl.append(deepcopy(modelo))
                fila = tabla.rows[-1]
            else:
                fila = tabla.add_row()
            for celda, valor in zip(fila.cells, [alumno.trimestre_ingreso, alumno.matricula, alumno.last_name, alumno.second_last_name, alumno.first_name]):
                celda.text = str(valor or '')
    for parrafo in parrafos(documento):
        for marcador, valor in valores.items():
            paragraph_replace_text(parrafo, re.compile(re.escape('{' + marcador + '}')), valor)
    pendientes = set(re.findall(r'\{[^{}]+\}', '\n'.join(p.text for p in parrafos(documento))))
    if pendientes:
        raise ValidationError('Quedaron marcadores sin reemplazar: ' + ', '.join(sorted(pendientes)))
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()
